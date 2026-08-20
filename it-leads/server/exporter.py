from __future__ import annotations

import csv
import io
import json
import re
from datetime import datetime, timezone
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

LEAD_COLUMNS = [
    ("business_name", "Business Name"),
    ("category", "Category"),
    ("types", "All Types"),
    ("phone", "Phone"),
    ("international_phone", "International Phone"),
    ("website", "Website"),
    ("email", "Email"),
    ("address", "Full Address"),
    ("city", "City"),
    ("state", "State"),
    ("postal_code", "Postal Code"),
    ("country", "Country"),
    ("rating", "Rating"),
    ("review_count", "Review Count"),
    ("business_status", "Status"),
    ("opening_hours", "Opening Hours"),
    ("latitude", "Latitude"),
    ("longitude", "Longitude"),
    ("google_maps_url", "Google Maps URL"),
    ("place_id", "Place ID"),
    ("description", "Description"),
]

DOWNLOAD_FORMATS = [
    {"id": "csv", "label": "CSV", "extension": "csv", "description": "Excel & Google Sheets compatible"},
    {"id": "xlsx", "label": "Excel", "extension": "xlsx", "description": "Formatted .xlsx spreadsheet"},
    {"id": "json", "label": "JSON", "extension": "json", "description": "Structured data for apps & CRM"},
    {"id": "txt", "label": "Text", "extension": "txt", "description": "Readable plain-text list"},
]

MEDIA_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "json": "application/json; charset=utf-8",
    "txt": "text/plain; charset=utf-8",
}


def slugify(value: str, max_len: int = 40) -> str:
    slug = re.sub(r"[^\w\s-]", "", value.strip().lower())
    slug = re.sub(r"[\s_-]+", "-", slug).strip("-")
    return slug[:max_len] or "leads"


def build_filename(
    *,
    location: str = "",
    category: str = "",
    format_id: str = "csv",
    custom_name: str = "",
) -> str:
    if custom_name.strip():
        base = slugify(custom_name.strip(), 60)
    else:
        parts = ["it-leads"]
        if location.strip():
            parts.append(slugify(location))
        if category.strip():
            parts.append(slugify(category))
        base = "-".join(parts)

    extension = next((item["extension"] for item in DOWNLOAD_FORMATS if item["id"] == format_id), format_id)
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d")
    return f"{base}-{timestamp}.{extension}"


def _row_values(lead: dict[str, Any], columns: list[tuple[str, str]] | None = None) -> list[Any]:
    cols = columns or LEAD_COLUMNS
    values: list[Any] = []
    for key, _ in cols:
        value = lead.get(key, "")
        values.append("" if value is None else value)
    return values


def leads_to_csv(leads: list[dict[str, Any]], columns: list[tuple[str, str]] | None = None) -> str:
    cols = columns or LEAD_COLUMNS
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow([label for _, label in cols])
    for lead in leads:
        writer.writerow(_row_values(lead, cols))
    return buffer.getvalue()


def leads_to_json(leads: list[dict[str, Any]], *, metadata: dict[str, Any] | None = None) -> str:
    payload: dict[str, Any] = {
        "exported_at": datetime.now(timezone.utc).isoformat(),
        "total": len(leads),
        "leads": leads,
    }
    if metadata:
        payload["search"] = metadata
    return json.dumps(payload, indent=2, ensure_ascii=False)


def leads_to_txt(leads: list[dict[str, Any]], *, metadata: dict[str, Any] | None = None) -> str:
    lines: list[str] = [
        "IT MAPS LEAD GENERATOR — EXPORT",
        f"Exported: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}",
        f"Total leads: {len(leads)}",
    ]

    if metadata:
        if metadata.get("location_used"):
            lines.append(f"Location: {metadata['location_used']}")
        if metadata.get("query_used"):
            lines.append(f"Query: {metadata['query_used']}")

    lines.append("=" * 60)

    for index, lead in enumerate(leads, start=1):
        lines.extend(
            [
                "",
                f"#{index} — {lead.get('business_name') or 'Unknown Business'}",
                f"Category: {lead.get('category') or '—'}",
                f"Phone: {lead.get('phone') or '—'}",
                f"Website: {lead.get('website') or '—'}",
                f"Address: {lead.get('address') or '—'}",
                f"City: {lead.get('city') or '—'}",
                f"State: {lead.get('state') or '—'}",
                f"Rating: {lead.get('rating') or '—'} ({lead.get('review_count') or 0} reviews)",
                f"Status: {lead.get('business_status') or '—'}",
                f"Hours: {lead.get('opening_hours') or '—'}",
                f"Google Maps: {lead.get('google_maps_url') or '—'}",
                "-" * 40,
            ]
        )

    return "\n".join(lines)


def leads_to_xlsx(leads: list[dict[str, Any]], columns: list[tuple[str, str]] | None = None) -> bytes:
    cols = columns or LEAD_COLUMNS
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Leads"

    header_fill = PatternFill("solid", fgColor="2563EB")
    header_font = Font(color="FFFFFF", bold=True)

    for col_index, (_, label) in enumerate(cols, start=1):
        cell = sheet.cell(row=1, column=col_index, value=label)
        cell.fill = header_fill
        cell.font = header_font

    for row_index, lead in enumerate(leads, start=2):
        for col_index, value in enumerate(_row_values(lead, cols), start=1):
            sheet.cell(row=row_index, column=col_index, value=value)

    for col_index, (_, label) in enumerate(cols, start=1):
        letter = get_column_letter(col_index)
        width = min(max(len(label) + 4, 12), 40)
        sheet.column_dimensions[letter].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def export_leads(
    leads: list[dict[str, Any]],
    format_id: str,
    *,
    columns: list[tuple[str, str]] | None = None,
    metadata: dict[str, Any] | None = None,
) -> tuple[bytes | str, str]:
    if format_id == "csv":
        return leads_to_csv(leads, columns), MEDIA_TYPES["csv"]
    if format_id == "json":
        return leads_to_json(leads, metadata=metadata), MEDIA_TYPES["json"]
    if format_id == "txt":
        return leads_to_txt(leads, metadata=metadata), MEDIA_TYPES["txt"]
    if format_id == "xlsx":
        return leads_to_xlsx(leads, columns), MEDIA_TYPES["xlsx"]
    raise ValueError(f"Unsupported format: {format_id}")
